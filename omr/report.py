"""Excel export of results."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
HEAD_FONT = Font(bold=True, color="FFFFFF")
WARN_FILL = PatternFill("solid", fgColor="FFE699")


def _sheet(wb, title, header, rows, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title[:31]
    ws.append(header)
    for c in ws[1]:
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    for i, h in enumerate(header, 1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:300] if i - 1 < len(r)])
        ws.column_dimensions[get_column_letter(i)].width = min(max(8, width + 2), 40)
    return ws


def build_workbook(exam: dict, payload: dict, problem_rows: list[dict]) -> bytes:
    wb = Workbook()
    sections = []
    for r in payload["rows"]:
        for s in r["sections"]:
            if s not in sections:
                sections.append(s)
    head = ["Rank", "Roll No", "Name", "Batch", "Set", "Correct", "Wrong", "Blank", "Multiple marked",
            "Score", "Max", "Percent"] + [f"{s} score" for s in sections]
    rows = []
    for r in payload["rows"]:
        rows.append([r["rank"], r["roll"], r["name"], r["batch"], r["set"], r["correct"], r["wrong"], r["blank"],
                     r["multi"], r["score"], r["max"], r["percent"]] + [r["sections"].get(s, {}).get("score", "") for s in sections])
    ws = _sheet(wb, "Results", head, rows, first=True)
    for i, r in enumerate(payload["rows"], 2):
        if r["needs_review"]:
            for c in ws[i]:
                c.fill = WARN_FILL

    _sheet(wb, "Batch summary", ["Batch", "Students", "Average", "Highest", "Lowest"],
           [[b["batch"], b["students"], b["average"], b["highest"], b["lowest"]] for b in payload["batches"]])

    labels = sorted({l for st in payload["stats"] for l in st["dist"]})
    qrows = []
    for st in payload["stats"]:
        qrows.append([st["set"], st["q"], st["key"], st["pct_correct"]] + [st["dist"].get(l, 0) for l in labels]
                     + [st["blank"], st["multi"], st["top"], "CHECK KEY" if st["suspicious"] else ""])
    ws = _sheet(wb, "Question analysis", ["Set", "Q", "Key", "% correct"] + [f"Marked {l}" for l in labels]
                + ["Blank", "Multiple", "Most chosen", "Note"], qrows)
    for i, st in enumerate(payload["stats"], 2):
        if st["suspicious"]:
            for c in ws[i]:
                c.fill = WARN_FILL

    krows = [[code, i + 1, tok] for code, key in exam["keys"].items() for i, tok in enumerate(key)]
    _sheet(wb, "Answer key", ["Set", "Q", "Correct option"], krows)

    prow = [[r["roll"] or "?", r["batch"], r["set"], r["file"], "; ".join(r["issues"] + r["warnings"])] for r in problem_rows]
    _sheet(wb, "Needs attention", ["Roll No", "Batch", "Set", "File", "Problem"], prow)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
